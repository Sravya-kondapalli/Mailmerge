import streamlit as st
import pandas as pd
from docx import Document
import io
import zipfile
import os

st.title("📧 Automated Certificate Generator")
st.write("Upload your recipient data file to generate and download your certificates instantly.")

# Only upload the recipient file now
uploaded_file = st.file_uploader("Upload recipients file (CSV or Excel)", type=["csv", "xlsx"])

# Path to the template in your repository
template_path = "certificate_template.docx"

if uploaded_file is not None:
    if not os.path.exists(template_path):
        st.error(f"Error: '{template_path}' was not found in your GitHub repository.")
    else:
        # Read recipient file
        if uploaded_file.name.endswith('.csv'):
            df = pd.read_csv(uploaded_file)
        else:
            df = pd.read_excel(uploaded_file)
        
        # Generate certificates in-memory using the repository's template
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, "w") as zip_file:
            for index, row in df.iterrows():
                doc = Document(template_path)
                
                # Replace placeholders in paragraphs
                for p in doc.paragraphs:
                    for key, val in row.items():
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in p.text:
                            p.text = p.text.replace(placeholder, str(val))
                
                # Replace placeholders in tables
                for table in doc.tables:
                    for row_t in table.rows:
                        for cell in row_t.cells:
                            for key, val in row.items():
                                placeholder = f"{{{{{key}}}}}"
                                if placeholder in cell.text:
                                    cell.text = cell.text.replace(placeholder, str(val))
                
                doc_io = io.BytesIO()
                doc.save(doc_io)
                doc_io.seek(0)
                
                file_name = f"Certificate_{row.get('Name', index)}.docx"
                zip_file.writestr(file_name, doc_io.read())
        
        zip_buffer.seek(0)
        
        st.success("Certificates generated successfully!")
        
        # Direct download button
        st.download_button(
            label="Download All Certificates (ZIP)",
            data=zip_buffer,
            file_name="Generated_Certificates.zip",
            mime="application/zip"
        )
